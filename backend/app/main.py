from __future__ import annotations

from contextlib import asynccontextmanager

import asyncio
import csv
from datetime import datetime, timezone
from io import BytesIO, StringIO
import logging
import math
import re
import sqlite3
import threading
import time
from typing import Annotated, AsyncIterator, cast
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError

from fastapi import Depends, FastAPI, File, Form, Header, HTTPException, Query, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from fastapi.exceptions import RequestValidationError

from .database import Database, IdempotencyConflictError, TaskStateConflictError
from .error_handling import public_error, validation_error_detail
from .schemas import (
    AliasCandidate, AliasReviewRequest, AliasReviewResult, BusinessMetrics,
    CatalogImportResult, CatalogPublishResult,
    Customer, Draft, DraftUpdate, HealthResponse, Order, Product, ProductCreate,
    ProductUpdate, RecognitionRequest,
    RecognitionTask,
    WorkflowCheckpoint,
    SkuIndexStatus, SkuVersion, SkuVersionPublishResult,
    CorrectionEvent, CustomerPreference,
    MetricsOverview, ModelMetrics, LlmFailureRecord,
    AgentMessageRequest, AgentResumeRequest, AgentSession, AgentSessionCreate,
)
from .settings import Settings
from .retrieval import SkuRetriever
from .input_normalizer import (
    MEDIA_FILE_LIMIT,
    STRUCTURED_FILE_LIMIT,
    normalize_uploaded_order_details,
)
from .workflow import OrderWorkflow, validate_edited_items
from .task_runner import RecognitionTaskRunner
from .task_events import RecognitionTaskEventBroker
from .agent_runtime.errors import (
    AgentIdempotencyConflict,
    AgentSessionNotFound,
    AgentStateConflict,
    ToolArgumentsInvalid,
    ToolExecutionFailed,
    ToolNotFound,
    ToolPermissionDenied,
)
from .agent_runtime.loop import AgentLoop
from .agent_runtime.gateway import DraftReviewLock


def get_settings(request: Request) -> Settings:
    return cast(Settings, request.app.state.settings)


def get_database(request: Request) -> Database:
    return cast(Database, request.app.state.database)


def get_retriever(request: Request) -> SkuRetriever:
    return cast(SkuRetriever, request.app.state.retriever)


def get_workflow(request: Request) -> OrderWorkflow:
    return cast(OrderWorkflow, request.app.state.workflow)


def get_task_runner(request: Request) -> RecognitionTaskRunner:
    return cast(RecognitionTaskRunner, request.app.state.task_runner)


def get_task_event_broker(request: Request) -> RecognitionTaskEventBroker:
    return cast(
        RecognitionTaskEventBroker, request.app.state.task_event_broker
    )


def get_agent_loop(request: Request) -> AgentLoop:
    return cast(AgentLoop, request.app.state.agent_loop)


def get_draft_review_lock(request: Request) -> DraftReviewLock:
    return cast(DraftReviewLock, request.app.state.draft_review_lock)


SettingsDep = Annotated[Settings, Depends(get_settings)]
DatabaseDep = Annotated[Database, Depends(get_database)]
SkuRetrieverDep = Annotated[SkuRetriever, Depends(get_retriever)]
OrderWorkflowDep = Annotated[OrderWorkflow, Depends(get_workflow)]
RecognitionTaskRunnerDep = Annotated[
    RecognitionTaskRunner, Depends(get_task_runner)
]
RecognitionTaskEventBrokerDep = Annotated[
    RecognitionTaskEventBroker, Depends(get_task_event_broker)
]
AgentLoopDep = Annotated[AgentLoop, Depends(get_agent_loop)]
DraftReviewLockDep = Annotated[DraftReviewLock, Depends(get_draft_review_lock)]

CATALOG_FILE_LIMIT = 2 * 1024 * 1024
CATALOG_ROW_LIMIT = 5000
XLSX_UNCOMPRESSED_LIMIT = 25 * 1024 * 1024
UPLOAD_CHUNK_SIZE = 64 * 1024
MULTIPART_ENVELOPE_ALLOWANCE = 64 * 1024
INDEX_SYNC_ERROR = "商品目录已保存，但搜索索引同步失败；请稍后手动同步"
SSE_HEARTBEAT_SECONDS = 15.0
TERMINAL_TASK_STATUSES = {"succeeded", "failed", "cancelled"}


def configure_agent_logging() -> logging.Logger:
    """Ensure local Agent logs remain visible under Uvicorn and debugpy."""
    # HTTP responses are logged selectively by the application middleware below.
    # Disabling Uvicorn's access logger avoids a second line for every request.
    logging.getLogger("uvicorn.access").disabled = True
    logger = logging.getLogger("agent.workflow")
    logger.setLevel(logging.INFO)
    logger.propagate = False

    if not any(handler.get_name() == "agent-console" for handler in logger.handlers):
        handler = logging.StreamHandler()
        handler.set_name("agent-console")
        handler.setFormatter(
            logging.Formatter("%(asctime)s %(levelname)s %(name)s %(message)s")
        )
        logger.addHandler(handler)

    return logger


def parse_catalog_upload(filename: str, content: bytes) -> list[dict]:
    if len(content) > CATALOG_FILE_LIMIT:
        raise ValueError("商品目录文件不能超过 2MB")
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "csv":
        reader = csv.DictReader(StringIO(content.decode("utf-8-sig")))
        rows = []
        for row_number, row in enumerate(reader, start=1):
            if row_number > CATALOG_ROW_LIMIT:
                raise ValueError(f"商品目录最多支持 {CATALOG_ROW_LIMIT} 行")
            rows.append(row)
    elif suffix == "xlsx":
        try:
            with ZipFile(BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > XLSX_UNCOMPRESSED_LIMIT:
                    raise ValueError("XLSX 商品目录解压后内容过大")
        except BadZipFile as exc:
            raise ValueError("XLSX 商品目录无法解析") from exc
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except (
            BadZipFile,
            EOFError,
            InvalidFileException,
            KeyError,
            OSError,
            ParseError,
            TypeError,
            ValueError,
        ) as exc:
            raise ValueError("XLSX 商品目录无法解析") from exc
        try:
            sheet = workbook.active
            values = []
            for row_number, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_number > CATALOG_ROW_LIMIT:
                    raise ValueError(f"商品目录最多支持 {CATALOG_ROW_LIMIT} 行")
                values.append(row)
        finally:
            workbook.close()
        if not values:
            rows = []
        else:
            headers = [str(value or "").strip().lower() for value in values[0]]
            rows = [dict(zip(headers, row)) for row in values[1:] if any(row)]
    else:
        raise ValueError("商品目录仅支持 CSV 或 XLSX")
    normalized_headers = {
        str(header or "").strip().lower() for header in (rows[0] if rows else {})
    }
    if rows:
        rows = [
            {str(key or "").strip().lower(): value for key, value in row.items()}
            for row in rows
        ]
    required = {"sku", "name", "unit", "unit_price"}
    if rows and not required.issubset(normalized_headers):
        raise ValueError("商品目录缺少 sku、name、unit 或 unit_price 列")
    def normalize_row(entry: tuple[int, dict]) -> dict:
        row_number, row = entry
        aliases_text = str(row.get("aliases") or "")
        active_text = str(row.get("active") if row.get("active") is not None else "1")
        raw_price = row.get("unit_price")
        if raw_price is None or str(raw_price).strip() == "":
            raise ValueError(f"商品目录第 {row_number} 行缺少单价")
        try:
            price = float(raw_price)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"商品目录第 {row_number} 行单价无效") from exc
        if not math.isfinite(price) or price < 0 or price > 1_000_000:
            raise ValueError(f"商品目录第 {row_number} 行单价无效")
        try:
            return ProductCreate(
                sku=str(row.get("sku") or "").strip(),
                name=str(row.get("name") or "").strip(),
                aliases=[
                    value.strip()
                    for value in re.split(r"[,;，；、|]", aliases_text)
                    if value.strip()
                ],
                unit=str(row.get("unit") or "").strip(),
                unit_price=price,
                active=active_text.strip().lower()
                not in {"0", "false", "no", "否", "停用"},
                source_type="import",
            ).model_dump()
        except ValidationError as exc:
            raise ValueError(f"商品目录第 {row_number} 行字段无效") from exc

    normalized_rows = list(map(normalize_row, enumerate(rows, start=2)))
    products_by_sku = {product["sku"]: product for product in normalized_rows}
    products_by_name = {product["name"]: product for product in normalized_rows}
    if len(products_by_sku) != len(normalized_rows):
        raise ValueError("导入文件包含重复 SKU")
    if len(products_by_name) != len(normalized_rows):
        raise ValueError("导入文件包含重复商品名称")
    return list(products_by_sku.values())


async def read_upload_limited(
    file: UploadFile, limit: int, too_large_message: str
) -> bytes:
    chunks: list[bytes] = []
    size = 0
    try:
        while chunk := await file.read(UPLOAD_CHUNK_SIZE):
            size += len(chunk)
            if size > limit:
                raise ValueError(too_large_message)
            chunks.append(chunk)
        return b"".join(chunks)
    finally:
        await file.close()


def order_upload_limit(filename: str) -> tuple[int, str]:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix in {"jpg", "jpeg", "png", "pdf"}:
        return MEDIA_FILE_LIMIT, "图片或 PDF 文件不能超过 8MB"
    return STRUCTURED_FILE_LIMIT, "TXT、CSV 或 XLSX 文件不能超过 2MB"


def safe_index_sync(
    retriever: SkuRetriever, site_id: int, sku_version: int
) -> dict:
    try:
        return retriever.sync_catalog(site_id, sku_version)
    except Exception as exc:
        logging.getLogger("agent.workflow").warning(
            "RAG action=sync_sku_index status=failed site_id=%d sku_version=%d reason=%s",
            site_id, sku_version, type(exc).__name__,
        )
        return {
            "site_id": site_id,
            "sku_version": sku_version,
            "indexed_count": 0,
            "backend": retriever.backend,
            "index_synced": False,
            "index_error": INDEX_SYNC_ERROR,
        }


def validate_recognition_context(
    database: Database, customer_id: int | None, site_id: int
) -> None:
    database.get_active_sku_version(site_id)
    if customer_id is not None and not database.customer_exists(customer_id):
        raise ValueError("所选客户不存在")


def _task_sse_event(task: dict) -> str:
    payload = RecognitionTask.model_validate(task).model_dump_json()
    return f"event: task\nid: {task['updated_at']}\ndata: {payload}\n\n"


async def recognition_task_event_stream(
    task_id: str,
    request: Request,
    database: Database,
    event_broker: RecognitionTaskEventBroker,
    heartbeat_seconds: float = SSE_HEARTBEAT_SECONDS,
) -> AsyncIterator[str]:
    subscription = event_broker.subscribe(task_id)
    try:
        current = await asyncio.to_thread(
            database.get_recognition_task, task_id
        )
        if current is None:
            return
        yield _task_sse_event(current)
        if current["status"] in TERMINAL_TASK_STATUSES:
            return
        while True:
            if await request.is_disconnected():
                return
            try:
                await subscription.get(heartbeat_seconds)
            except TimeoutError:
                if await request.is_disconnected():
                    return
                yield ": heartbeat\n\n"
                continue
            current = await asyncio.to_thread(
                database.get_recognition_task, task_id
            )
            if current is None:
                return
            yield _task_sse_event(current)
            if current["status"] in TERMINAL_TASK_STATUSES:
                return
    finally:
        subscription.close()


def create_app(database_path: str | None = None) -> FastAPI:
    agent_logger = configure_agent_logging()
    settings = Settings.from_env(database_path)
    database = Database(settings.database_path)

    @asynccontextmanager
    async def lifespan(app: FastAPI):
        database.initialize()
        app.state.settings = settings
        app.state.database = database
        app.state.retriever = SkuRetriever(database, settings)
        app.state.workflow = OrderWorkflow(database, settings, app.state.retriever)
        app.state.task_event_broker = RecognitionTaskEventBroker()
        app.state.task_runner = RecognitionTaskRunner(
            database,
            app.state.workflow,
            settings,
            app.state.task_event_broker,
        )
        app.state.draft_review_lock = threading.RLock()
        app.state.agent_loop = AgentLoop(
            database,
            app.state.workflow,
            app.state.task_runner,
            settings,
            app.state.draft_review_lock,
        )
        app.state.agent_loop.initialize()
        if settings.elasticsearch_url:
            try:
                current_version = database.get_active_sku_version(1)
                app.state.retriever.sync_catalog(1, current_version)
            except Exception as exc:
                agent_logger.warning(
                    "RAG action=startup_sync status=fallback reason=%s",
                    type(exc).__name__,
                )
        app.state.task_runner.start()
        try:
            yield
        finally:
            app.state.task_runner.close()
            app.state.workflow.close()
            app.state.retriever.close()

    app = FastAPI(
        title="Local AI Order Agent",
        version="0.1.0",
        lifespan=lifespan,
    )
    app.add_middleware(
        CORSMiddleware,
        allow_origins=["http://localhost:3000", "http://localhost:8080"],
        allow_methods=["*"],
        allow_headers=["*"],
    )

    @app.exception_handler(RequestValidationError)
    async def request_validation_error(
        request: Request, error: RequestValidationError
    ) -> JSONResponse:
        return JSONResponse(
            status_code=422,
            content={"detail": validation_error_detail(error.errors())},
        )

    @app.middleware("http")
    async def log_http_request(request: Request, call_next):
        upload_request_limits = {
            "/api/catalog/import": CATALOG_FILE_LIMIT + MULTIPART_ENVELOPE_ALLOWANCE,
            "/api/drafts/recognize-file": MEDIA_FILE_LIMIT + MULTIPART_ENVELOPE_ALLOWANCE,
            "/api/recognition-tasks/file": MEDIA_FILE_LIMIT + MULTIPART_ENVELOPE_ALLOWANCE,
        }
        request_limit = upload_request_limits.get(request.url.path)
        content_length = request.headers.get("content-length", "")
        if request_limit is not None and content_length.isdigit():
            if int(content_length) > request_limit:
                return JSONResponse(
                    status_code=413,
                    content={"detail": "上传文件过大"},
                )
        started = time.perf_counter()
        try:
            response = await call_next(request)
        except Exception as exc:
            agent_logger.error(
                "HTTP method=%s path=%s status=failed reason=%s",
                request.method, request.url.path, type(exc).__name__,
            )
            raise
        duration_ms = round((time.perf_counter() - started) * 1000)
        if response.status_code != 200:
            level = (
                logging.ERROR if response.status_code >= 500
                else logging.WARNING if response.status_code >= 400
                else logging.INFO
            )
            agent_logger.log(
                level,
                "HTTP method=%s path=%s status=%d duration_ms=%d",
                request.method,
                request.url.path,
                response.status_code,
                duration_ms,
            )
        return response

    @app.get("/api/health", response_model=HealthResponse)
    def health(current_settings: SettingsDep) -> HealthResponse:
        return HealthResponse(
            status="ok",
            ai_provider=current_settings.ai_provider,
            database="sqlite",
        )

    @app.post(
        "/api/agent/sessions", response_model=AgentSession, status_code=201
    )
    def create_agent_session(
        payload: AgentSessionCreate, agent_loop: AgentLoopDep
    ) -> dict:
        try:
            return agent_loop.create_session(payload.site_id, payload.customer_id)
        except LookupError as exc:
            raise HTTPException(status_code=404, detail="站点不存在") from exc
        except ValueError as exc:
            raise HTTPException(
                status_code=422, detail=public_error(exc, "会话参数无效")
            ) from exc

    @app.post(
        "/api/agent/sessions/{session_id}/messages",
        response_model=AgentSession,
    )
    def send_agent_message(
        session_id: str,
        payload: AgentMessageRequest,
        agent_loop: AgentLoopDep,
        idempotency_key: str | None = Header(
            default=None, alias="Idempotency-Key", max_length=200
        ),
    ) -> dict:
        try:
            return agent_loop.handle_message(
                session_id, payload.content, idempotency_key=idempotency_key
            )
        except AgentSessionNotFound as exc:
            raise HTTPException(status_code=404, detail=str(exc)) from exc
        except (AgentStateConflict, AgentIdempotencyConflict) as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        except (
            ToolArgumentsInvalid, ToolExecutionFailed, ToolNotFound,
            ToolPermissionDenied,
        ) as exc:
            raise HTTPException(
                status_code=422, detail=public_error(exc, "Agent 工具执行失败")
            ) from exc

    @app.get(
        "/api/agent/sessions/{session_id}", response_model=AgentSession
    )
    def get_agent_session(
        session_id: str,
        agent_loop: AgentLoopDep,
        after_sequence: int | None = Query(default=None, ge=0),
        limit: int = Query(default=100, ge=1, le=200),
    ) -> dict:
        try:
            return agent_loop.get_session(
                session_id, after_sequence=after_sequence, event_limit=limit
            )
        except AgentSessionNotFound as exc:
            raise HTTPException(status_code=404, detail=str(exc)) from exc
        except AgentStateConflict as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc

    @app.post(
        "/api/agent/sessions/{session_id}/resume",
        response_model=AgentSession,
    )
    def resume_agent_session(
        session_id: str,
        payload: AgentResumeRequest,
        agent_loop: AgentLoopDep,
        idempotency_key: str | None = Header(
            default=None, alias="Idempotency-Key", max_length=200
        ),
    ) -> dict:
        try:
            return agent_loop.resume(
                session_id, payload.approved, idempotency_key=idempotency_key
            )
        except AgentSessionNotFound as exc:
            raise HTTPException(status_code=404, detail=str(exc)) from exc
        except (AgentStateConflict, AgentIdempotencyConflict) as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        except (
            ToolArgumentsInvalid, ToolExecutionFailed, ToolNotFound,
            ToolPermissionDenied,
        ) as exc:
            raise HTTPException(
                status_code=422, detail=public_error(exc, "Agent 恢复执行失败")
            ) from exc

    @app.get("/api/catalog/customers", response_model=list[Customer])
    def list_customers(database: DatabaseDep) -> list[dict]:
        return database.list_customers()

    @app.get("/api/catalog/products", response_model=list[Product])
    def list_products(database: DatabaseDep) -> list[dict]:
        return database.list_products()

    @app.post("/api/catalog/products", response_model=CatalogPublishResult)
    def create_product(
        payload: ProductCreate,
        database: DatabaseDep,
        retriever: SkuRetrieverDep,
        site_id: int = Query(..., ge=1),
    ) -> CatalogPublishResult:
        try:
            product, version = database.create_product(
                payload.model_dump(), site_id
            )
            index_status = safe_index_sync(
                retriever,
                site_id, version["version"]
            )
        except sqlite3.IntegrityError as exc:
            raise HTTPException(
                status_code=409, detail="商品目录存在冲突，请刷新后重试"
            ) from exc
        except sqlite3.OperationalError as exc:
            raise HTTPException(
                status_code=409, detail="商品目录正在更新，请稍后重试"
            ) from exc
        except LookupError as exc:
            raise HTTPException(status_code=404, detail="站点不存在") from exc
        except ValueError as exc:
            raise HTTPException(
                status_code=422, detail=public_error(exc, "商品信息无效")
            ) from exc
        return CatalogPublishResult(
            product=product,
            version=version,
            indexed_count=index_status["indexed_count"],
            backend=index_status["backend"],
            index_synced=index_status["index_synced"],
            index_error=index_status["index_error"],
        )

    @app.put(
        "/api/catalog/products/{product_id}", response_model=CatalogPublishResult
    )
    def update_product(
        product_id: int,
        payload: ProductUpdate,
        database: DatabaseDep,
        retriever: SkuRetrieverDep,
        site_id: int = Query(..., ge=1),
    ) -> CatalogPublishResult:
        try:
            product, version = database.update_product(
                product_id, payload.model_dump(), site_id
            )
            index_status = safe_index_sync(
                retriever,
                site_id, version["version"]
            )
        except LookupError as exc:
            raise HTTPException(
                status_code=404, detail=public_error(exc, "商品不存在")
            ) from exc
        except sqlite3.IntegrityError as exc:
            raise HTTPException(
                status_code=409, detail="商品目录存在冲突，请刷新后重试"
            ) from exc
        except sqlite3.OperationalError as exc:
            raise HTTPException(
                status_code=409, detail="商品目录正在更新，请稍后重试"
            ) from exc
        except ValueError as exc:
            raise HTTPException(
                status_code=422, detail=public_error(exc, "商品信息无效")
            ) from exc
        return CatalogPublishResult(
            product=product,
            version=version,
            indexed_count=index_status["indexed_count"],
            backend=index_status["backend"],
            index_synced=index_status["index_synced"],
            index_error=index_status["index_error"],
        )

    @app.post("/api/catalog/import", response_model=CatalogImportResult)
    async def import_catalog(
        database: DatabaseDep,
        retriever: SkuRetrieverDep,
        file: UploadFile = File(...),
        site_id: int = Form(..., ge=1),
    ) -> CatalogImportResult:
        try:
            content = await read_upload_limited(
                file, CATALOG_FILE_LIMIT, "商品目录文件不能超过 2MB"
            )
            products = await asyncio.to_thread(
                parse_catalog_upload, file.filename or "", content
            )
            result = await asyncio.to_thread(database.import_products, products, site_id)
            index_status = await asyncio.to_thread(
                safe_index_sync,
                retriever,
                site_id,
                result["version"]["version"],
            )
        except UnicodeDecodeError as exc:
            raise HTTPException(status_code=422, detail="CSV 文件必须使用 UTF-8 编码") from exc
        except (TypeError, ValueError) as exc:
            raise HTTPException(
                status_code=422,
                detail=public_error(exc, "商品目录无法解析"),
            ) from exc
        except LookupError as exc:
            raise HTTPException(status_code=404, detail="站点不存在") from exc
        except sqlite3.IntegrityError as exc:
            raise HTTPException(
                status_code=409, detail="商品目录存在冲突，请刷新后重试"
            ) from exc
        except sqlite3.OperationalError as exc:
            raise HTTPException(
                status_code=409, detail="商品目录正在更新，请稍后重试"
            ) from exc
        return CatalogImportResult(
            imported_count=result["imported_count"],
            version=result["version"],
            indexed_count=index_status["indexed_count"],
            backend=index_status["backend"],
            index_synced=index_status["index_synced"],
            index_error=index_status["index_error"],
        )

    @app.get(
        "/api/catalog/alias-candidates", response_model=list[AliasCandidate]
    )
    def list_alias_candidates(
        database: DatabaseDep,
        status: str | None = Query(
            default=None, pattern="^(pending|approved|rejected)$"
        ),
        limit: int = Query(default=100, ge=1, le=500),
    ) -> list[dict]:
        return database.list_alias_candidates(status, limit)

    @app.post(
        "/api/catalog/alias-candidates/{candidate_id}/review",
        response_model=AliasReviewResult,
    )
    def review_alias_candidate(
        candidate_id: int,
        payload: AliasReviewRequest,
        database: DatabaseDep,
        retriever: SkuRetrieverDep,
    ) -> AliasReviewResult:
        try:
            candidate, version = database.review_alias_candidate(
                candidate_id, payload.decision
            )
        except LookupError as exc:
            raise HTTPException(
                status_code=404, detail=public_error(exc, "别名候选不存在")
            ) from exc
        except ValueError as exc:
            raise HTTPException(
                status_code=409, detail=public_error(exc, "别名候选状态冲突")
            ) from exc
        if version is None:
            return AliasReviewResult(candidate=candidate, version=None)
        index_status = safe_index_sync(
            retriever,
            candidate["site_id"], version["version"]
        )
        return AliasReviewResult(
            candidate=candidate,
            version=version,
            indexed_count=index_status["indexed_count"],
            backend=index_status["backend"],
            index_synced=index_status["index_synced"],
            index_error=index_status["index_error"],
        )

    @app.get("/api/metrics/business", response_model=BusinessMetrics)
    def business_metrics(database: DatabaseDep, site_id: int = 1) -> dict:
        try:
            return database.get_business_metrics(site_id)
        except LookupError as exc:
            raise HTTPException(status_code=404, detail="站点不存在") from exc

    @app.get("/api/catalog/versions/current", response_model=SkuVersion)
    def current_sku_version(database: DatabaseDep, site_id: int = 1) -> dict:
        try:
            version = database.get_active_sku_version(site_id)
        except LookupError as exc:
            raise HTTPException(status_code=404, detail="站点不存在") from exc
        result = database.get_sku_version(site_id, version)
        if result is None:
            raise HTTPException(status_code=404, detail="SKU 版本不存在")
        return result

    @app.post("/api/catalog/versions", response_model=SkuVersionPublishResult)
    def create_sku_version(
        database: DatabaseDep,
        retriever: SkuRetrieverDep,
        site_id: int = Query(..., ge=1),
    ) -> SkuVersionPublishResult:
        try:
            version = database.create_sku_version(site_id)
        except LookupError as exc:
            raise HTTPException(status_code=404, detail="站点不存在") from exc
        except (sqlite3.IntegrityError, sqlite3.OperationalError) as exc:
            raise HTTPException(
                status_code=409, detail="商品目录正在更新，请稍后重试"
            ) from exc
        index_status = safe_index_sync(retriever, site_id, version["version"])
        return SkuVersionPublishResult(
            version=version,
            indexed_count=index_status["indexed_count"],
            backend=index_status["backend"],
            index_synced=index_status["index_synced"],
            index_error=index_status["index_error"],
        )

    @app.post("/api/catalog/index/sync", response_model=SkuIndexStatus)
    def sync_sku_index(
        database: DatabaseDep,
        retriever: SkuRetrieverDep,
        site_id: int = Query(..., ge=1),
    ) -> dict:
        try:
            version = database.get_active_sku_version(site_id)
        except LookupError as exc:
            raise HTTPException(status_code=404, detail="站点不存在") from exc
        return safe_index_sync(retriever, site_id, version)

    @app.post("/api/drafts/recognize", response_model=Draft)
    def recognize(
        payload: RecognitionRequest,
        workflow: OrderWorkflowDep,
        database: DatabaseDep,
    ) -> dict:
        try:
            validate_recognition_context(
                database, payload.customer_id, payload.site_id
            )
            return workflow.recognize(
                payload.text, payload.customer_id, payload.site_id
            )
        except LookupError as exc:
            raise HTTPException(status_code=404, detail="站点不存在") from exc
        except ValueError as exc:
            raise HTTPException(
                status_code=422, detail=public_error(exc, "订单内容无法识别")
            ) from exc

    @app.post(
        "/api/recognition-tasks", response_model=RecognitionTask, status_code=202
    )
    def create_recognition_task(
        payload: RecognitionRequest,
        task_runner: RecognitionTaskRunnerDep,
        database: DatabaseDep,
        idempotency_key: str | None = Header(
            default=None, alias="Idempotency-Key"
        ),
    ) -> dict:
        try:
            validate_recognition_context(
                database, payload.customer_id, payload.site_id
            )
            return task_runner.create_task(
                payload.text,
                payload.customer_id,
                payload.site_id,
                idempotency_key=idempotency_key,
            )
        except IdempotencyConflictError as exc:
            raise HTTPException(
                status_code=409, detail=public_error(exc, "识别任务幂等键冲突")
            ) from exc
        except LookupError as exc:
            raise HTTPException(status_code=404, detail="站点不存在") from exc
        except ValueError as exc:
            raise HTTPException(
                status_code=422, detail=public_error(exc, "订单内容无法识别")
            ) from exc

    @app.get("/api/recognition-tasks", response_model=list[RecognitionTask])
    def list_recognition_tasks(
        database: DatabaseDep,
        task_status: str | None = Query(
            default=None,
            alias="status",
            pattern="^(pending|running|retrying|succeeded|failed|cancelled)$",
        ),
        limit: int = Query(default=50, ge=1, le=200),
    ) -> list[dict]:
        return database.list_recognition_tasks(
            status=task_status, limit=limit
        )

    @app.get(
        "/api/recognition-tasks/{task_id}", response_model=RecognitionTask
    )
    def get_recognition_task(task_id: str, database: DatabaseDep) -> dict:
        task = database.get_recognition_task(task_id)
        if task is None:
            raise HTTPException(status_code=404, detail="识别任务不存在")
        return task

    @app.get("/api/recognition-tasks/{task_id}/events")
    async def stream_recognition_task_events(
        task_id: str,
        request: Request,
        database: DatabaseDep,
        event_broker: RecognitionTaskEventBrokerDep,
    ) -> StreamingResponse:
        initial_task = await asyncio.to_thread(
            database.get_recognition_task, task_id
        )
        if initial_task is None:
            raise HTTPException(status_code=404, detail="识别任务不存在")

        return StreamingResponse(
            recognition_task_event_stream(
                task_id, request, database, event_broker
            ),
            media_type="text/event-stream",
            headers={
                "Cache-Control": "no-cache",
                "X-Accel-Buffering": "no",
                "Connection": "keep-alive",
            },
        )

    @app.post(
        "/api/recognition-tasks/{task_id}/retry",
        response_model=RecognitionTask,
    )
    def retry_recognition_task(
        task_id: str,
        database: DatabaseDep,
        task_runner: RecognitionTaskRunnerDep,
    ) -> dict:
        now = datetime.now(timezone.utc).isoformat(timespec="milliseconds")
        try:
            task = database.retry_recognition_task(task_id, now)
        except LookupError as exc:
            raise HTTPException(
                status_code=404, detail=public_error(exc, "识别任务不存在")
            ) from exc
        except TaskStateConflictError as exc:
            raise HTTPException(
                status_code=409, detail=public_error(exc, "识别任务状态冲突")
            ) from exc
        task_runner.wake()
        task_runner.publish(task)
        return task

    @app.post(
        "/api/recognition-tasks/{task_id}/cancel",
        response_model=RecognitionTask,
    )
    def cancel_recognition_task(
        task_id: str,
        database: DatabaseDep,
        task_runner: RecognitionTaskRunnerDep,
    ) -> dict:
        now = datetime.now(timezone.utc).isoformat(timespec="milliseconds")
        try:
            task = database.cancel_recognition_task(task_id, now)
        except LookupError as exc:
            raise HTTPException(
                status_code=404, detail=public_error(exc, "识别任务不存在")
            ) from exc
        task_runner.wake()
        task_runner.publish(task)
        return task

    @app.post("/api/drafts/recognize-file", response_model=Draft)
    async def recognize_file(
        workflow: OrderWorkflowDep,
        database: DatabaseDep,
        file: UploadFile = File(...),
        customer_id: int = Form(...),
        site_id: int = Form(..., ge=1),
    ) -> dict:
        filename = file.filename or ""
        try:
            limit, limit_message = order_upload_limit(filename)
            content = await read_upload_limited(file, limit, limit_message)
            normalized = await asyncio.to_thread(
                normalize_uploaded_order_details, filename, content
            )
            validate_recognition_context(database, customer_id, site_id)
            return await asyncio.to_thread(
                workflow.recognize,
                normalized.text,
                customer_id,
                site_id,
                normalized.input_type,
                filename,
                normalized.confidence,
                normalized.warnings,
            )
        except UnicodeDecodeError as exc:
            raise HTTPException(status_code=422, detail="上传文件必须使用 UTF-8 编码") from exc
        except LookupError as exc:
            raise HTTPException(status_code=404, detail="站点不存在") from exc
        except ValueError as exc:
            raise HTTPException(
                status_code=422,
                detail=public_error(exc, "上传文件无法识别"),
            ) from exc

    @app.post(
        "/api/recognition-tasks/file",
        response_model=RecognitionTask,
        status_code=202,
    )
    async def create_file_recognition_task(
        task_runner: RecognitionTaskRunnerDep,
        database: DatabaseDep,
        file: UploadFile = File(...),
        customer_id: int = Form(...),
        site_id: int = Form(..., ge=1),
        idempotency_key: str | None = Header(default=None, alias="Idempotency-Key"),
    ) -> dict:
        filename = file.filename or ""
        try:
            limit, limit_message = order_upload_limit(filename)
            content = await read_upload_limited(file, limit, limit_message)
            normalized = await asyncio.to_thread(
                normalize_uploaded_order_details, filename, content
            )
            validate_recognition_context(database, customer_id, site_id)
            return await asyncio.to_thread(
                task_runner.create_task,
                normalized.text,
                customer_id,
                site_id,
                idempotency_key,
                normalized.input_type,
                filename,
                normalized.confidence,
                normalized.warnings,
            )
        except IdempotencyConflictError as exc:
            raise HTTPException(
                status_code=409, detail=public_error(exc, "识别任务幂等键冲突")
            ) from exc
        except UnicodeDecodeError as exc:
            raise HTTPException(status_code=422, detail="上传文件必须使用 UTF-8 编码") from exc
        except LookupError as exc:
            raise HTTPException(status_code=404, detail="站点不存在") from exc
        except ValueError as exc:
            raise HTTPException(
                status_code=422,
                detail=public_error(exc, "上传文件无法识别"),
            ) from exc

    @app.get("/api/drafts", response_model=list[Draft])
    def list_drafts(
        database: DatabaseDep,
        status: str | None = Query(
            default=None, pattern="^(needs_review|confirmed)$"
        ),
        limit: int = Query(default=50, ge=1, le=200),
    ) -> list[dict]:
        return database.list_drafts(status=status, limit=limit)

    @app.get("/api/drafts/{draft_id}", response_model=Draft)
    def get_draft(draft_id: str, database: DatabaseDep) -> dict:
        draft = database.get_draft(draft_id)
        if not draft:
            raise HTTPException(status_code=404, detail="草稿不存在")
        return draft

    @app.get(
        "/api/drafts/{draft_id}/workflow", response_model=WorkflowCheckpoint
    )
    def get_workflow_checkpoint(
        draft_id: str,
        database: DatabaseDep,
        workflow: OrderWorkflowDep,
    ) -> dict:
        if not database.get_draft(draft_id):
            raise HTTPException(status_code=404, detail="草稿不存在")
        return workflow.checkpoint_status(draft_id)

    @app.get(
        "/api/drafts/{draft_id}/corrections", response_model=list[CorrectionEvent]
    )
    def list_draft_corrections(
        draft_id: str, database: DatabaseDep
    ) -> list[dict]:
        if not database.get_draft(draft_id):
            raise HTTPException(status_code=404, detail="草稿不存在")
        return database.list_corrections(draft_id)

    @app.get(
        "/api/customers/{customer_id}/preferences",
        response_model=list[CustomerPreference],
    )
    def list_customer_preferences(
        customer_id: int, database: DatabaseDep
    ) -> list[dict]:
        if not database.customer_exists(customer_id):
            raise HTTPException(status_code=404, detail="客户不存在")
        return database.list_customer_preferences(customer_id)

    @app.get("/api/metrics/overview", response_model=MetricsOverview)
    def metrics_overview(database: DatabaseDep) -> dict:
        return database.get_metrics_overview()

    @app.get("/api/metrics/models", response_model=list[ModelMetrics])
    def model_metrics(database: DatabaseDep) -> list[dict]:
        return database.get_model_metrics()

    @app.get("/api/metrics/failures", response_model=list[LlmFailureRecord])
    def model_failures(
        database: DatabaseDep,
        limit: int = Query(default=50, ge=1, le=200),
    ) -> list[dict]:
        return database.list_llm_failures(limit)

    @app.put("/api/drafts/{draft_id}", response_model=Draft)
    def update_draft(
        draft_id: str,
        payload: DraftUpdate,
        database: DatabaseDep,
        agent_loop: AgentLoopDep,
        draft_review_lock: DraftReviewLockDep,
    ) -> dict:
        with draft_review_lock:
            existing = database.get_draft(draft_id)
            if not existing:
                raise HTTPException(status_code=404, detail="草稿不存在")
            try:
                agent_loop.assert_draft_customer_change_allowed(
                    draft_id, payload.customer_id
                )
            except AgentStateConflict as exc:
                raise HTTPException(status_code=409, detail=str(exc)) from exc
            warnings = validate_edited_items(
                database,
                payload.customer_id,
                [item.model_dump() for item in payload.items],
                existing["site_id"],
                existing["sku_version"],
            )
            now = datetime.now(timezone.utc).isoformat(timespec="seconds")
            draft = database.update_draft(
                draft_id,
                payload.customer_id,
                [item.model_dump() for item in payload.items],
                warnings,
                now,
            )
            if not draft:
                raise HTTPException(
                    status_code=409, detail="草稿不存在或已经确认"
                )
            return draft

    @app.post("/api/drafts/{draft_id}/confirm", response_model=Order)
    def confirm_draft(
        draft_id: str,
        workflow: OrderWorkflowDep,
        draft_review_lock: DraftReviewLockDep,
        idempotency_key: str | None = Header(default=None, alias="Idempotency-Key"),
    ) -> dict:
        with draft_review_lock:
            try:
                order = workflow.resume(
                    draft_id, approved=True, idempotency_key=idempotency_key
                )
            except IdempotencyConflictError as exc:
                raise HTTPException(
                    status_code=409, detail=public_error(exc, "订单幂等键冲突")
                ) from exc
            except LookupError as exc:
                raise HTTPException(
                    status_code=404, detail=public_error(exc, "草稿不存在")
                ) from exc
            except ValueError as exc:
                raise HTTPException(
                    status_code=422, detail=public_error(exc, "订单无法确认")
                ) from exc
            except sqlite3.IntegrityError as exc:
                raise HTTPException(
                    status_code=409, detail="订单已确认或幂等键发生冲突"
                ) from exc
            return order

    @app.get("/api/orders", response_model=list[Order])
    def list_orders(database: DatabaseDep) -> list[dict]:
        return database.list_orders()

    @app.get("/api/orders/{order_id}", response_model=Order)
    def get_order(order_id: int, database: DatabaseDep) -> dict:
        order = database.get_order(order_id)
        if not order:
            raise HTTPException(status_code=404, detail="订单不存在")
        return order

    return app


app = create_app()
