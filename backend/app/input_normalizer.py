from __future__ import annotations

import csv
import io
import os
import re
import shutil
from dataclasses import dataclass, field
from itertools import groupby
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import pymupdf as fitz
import pytesseract
from openpyxl import load_workbook
from PIL import Image, UnidentifiedImageError
from pytesseract import Output, TesseractNotFoundError


REQUIRED_HEADERS = ("商品", "数量", "单位")
STRUCTURED_FILE_LIMIT = 2 * 1024 * 1024
MEDIA_FILE_LIMIT = 8 * 1024 * 1024
PDF_PAGE_LIMIT = 5
MAX_NORMALIZED_TEXT_LENGTH = 5000
MAX_ORDER_ROWS = 200
MAX_IMAGE_EDGE = 12_000
MAX_IMAGE_PIXELS = 25_000_000
XLSX_UNCOMPRESSED_LIMIT = 25 * 1024 * 1024
LOW_OCR_CONFIDENCE = 70.0
OCR_LANGUAGES = ("chi_sim", "eng")
PROJECT_TESSERACT = (
    Path(__file__).resolve().parents[2]
    / ".venv"
    / "tools"
    / "tesseract"
    / "tesseract.exe"
)


@dataclass(frozen=True)
class NormalizedInput:
    text: str
    input_type: str
    confidence: float | None = None
    warnings: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class OcrCandidate:
    language: str
    text: str
    confidence: float


def normalize_uploaded_order(filename: str, content: bytes) -> tuple[str, str]:
    """Backward-compatible wrapper used by existing callers and tests."""
    result = normalize_uploaded_order_details(filename, content)
    return result.text, result.input_type


def normalize_uploaded_order_details(filename: str, content: bytes) -> NormalizedInput:
    suffix = Path(filename).suffix.lower()
    if suffix in {".jpg", ".jpeg", ".png", ".pdf"}:
        if len(content) > MEDIA_FILE_LIMIT:
            raise ValueError("图片或 PDF 文件不能超过 8MB")
    elif len(content) > STRUCTURED_FILE_LIMIT:
        raise ValueError("TXT、CSV 或 XLSX 文件不能超过 2MB")

    if suffix == ".txt":
        text = content.decode("utf-8-sig").strip()
        if not text:
            raise ValueError("TXT 文件不能为空")
        return NormalizedInput(_bounded_text(text), "txt")
    if suffix == ".csv":
        rows = []
        for index, row in enumerate(
            csv.DictReader(io.StringIO(content.decode("utf-8-sig"))), start=1
        ):
            if index > MAX_ORDER_ROWS:
                raise ValueError(f"订单文件最多支持 {MAX_ORDER_ROWS} 行商品")
            rows.append(row)
        return NormalizedInput(_bounded_text(_rows_to_text(rows)), "csv")
    if suffix == ".xlsx":
        try:
            with ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > XLSX_UNCOMPRESSED_LIMIT:
                    raise ValueError("XLSX 文件解压后内容过大")
        except BadZipFile as exc:
            raise ValueError("XLSX 文件无效或无法读取") from exc
        try:
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=True
            )
        except Exception as exc:
            raise ValueError("XLSX 文件无效或无法读取") from exc
        try:
            sheet = workbook.active
            values = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index > MAX_ORDER_ROWS:
                    raise ValueError(f"订单文件最多支持 {MAX_ORDER_ROWS} 行商品")
                values.append(row)
        finally:
            workbook.close()
        if not values:
            raise ValueError("XLSX 文件不能为空")
        headers = [str(value).strip() if value is not None else "" for value in values[0]]
        rows = [dict(zip(headers, row)) for row in values[1:]]
        return NormalizedInput(_bounded_text(_rows_to_text(rows)), "xlsx")
    if suffix in {".jpg", ".jpeg", ".png"}:
        return _normalize_image(content)
    if suffix == ".pdf":
        return _normalize_pdf(content)
    raise ValueError("仅支持 UTF-8 TXT、CSV、XLSX、JPG、PNG 和 PDF 文件")


def _normalize_image(content: bytes) -> NormalizedInput:
    try:
        with Image.open(io.BytesIO(content)) as source:
            _validate_image_dimensions(source.width, source.height)
            image = source.convert("RGB")
    except (UnidentifiedImageError, OSError) as exc:
        raise ValueError("图片文件无效或无法读取") from exc
    text, confidence = _ocr_image(image)
    return NormalizedInput(text, "image_ocr", confidence, _ocr_warnings(confidence))


def _normalize_pdf(content: bytes) -> NormalizedInput:
    try:
        document = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise ValueError("PDF 文件无效或无法读取") from exc
    try:
        if document.page_count == 0:
            raise ValueError("PDF 文件没有页面")
        if document.page_count > PDF_PAGE_LIMIT:
            raise ValueError(f"PDF 最多支持 {PDF_PAGE_LIMIT} 页")
        texts: list[str] = []
        confidences: list[float] = []
        ocr_pages = 0
        text_pages = 0
        for page in document:
            page_text = page.get_text("text").strip()
            if page_text:
                texts.append(page_text)
                text_pages += 1
                continue
            render_width = int(page.rect.width * 2)
            render_height = int(page.rect.height * 2)
            _validate_image_dimensions(render_width, render_height)
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            ocr_text, confidence = _ocr_image(image)
            texts.append(ocr_text)
            confidences.append(confidence)
            ocr_pages += 1
        text = "\n".join(part for part in texts if part).strip()
        if not text:
            raise ValueError("PDF 中没有可识别的订单内容")
        text = _bounded_text(text)
        confidence = round(sum(confidences) / len(confidences), 2) if confidences else None
        input_type = (
            "pdf_text" if ocr_pages == 0 else "pdf_ocr" if text_pages == 0 else "pdf_mixed"
        )
        return NormalizedInput(text, input_type, confidence, _ocr_warnings(confidence))
    finally:
        document.close()


def _ocr_image(image: Image.Image) -> tuple[str, float]:
    tesseract_config = _configure_tesseract()
    try:
        candidates = tuple(
            filter(
                None,
                map(
                    lambda language: _ocr_candidate(image, language, tesseract_config),
                    OCR_LANGUAGES,
                ),
            )
        )
    except TesseractNotFoundError as exc:
        raise ValueError("本机未安装 Tesseract OCR；请安装后重启后端") from exc
    if not candidates:
        raise ValueError("图片中没有识别到订单内容")
    selected = max(candidates, key=_ocr_candidate_rank)
    return selected.text, selected.confidence


def _ocr_candidate(
    image: Image.Image, language: str, tesseract_config: str
) -> OcrCandidate | None:
    try:
        data = pytesseract.image_to_data(
            image,
            lang=language,
            config=tesseract_config,
            output_type=Output.DICT,
        )
    except TesseractNotFoundError:
        raise
    except pytesseract.TesseractError:
        return None
    texts = tuple(map(lambda value: str(value).strip(), data.get("text", [])))
    line_parts = [
        ((page, block, paragraph, line), text)
        for page, block, paragraph, line, text in zip(
            data.get("page_num", []),
            data.get("block_num", []),
            data.get("par_num", []),
            data.get("line_num", []),
            texts,
        )
        if text
    ]
    text = "\n".join(
        _join_ocr_line(tuple(part for _, part in group))
        for _, group in groupby(line_parts, key=lambda item: item[0])
    ).strip()
    if not text:
        return None
    scores = tuple(filter(lambda score: score >= 0, map(_confidence_value, data.get("conf", []))))
    confidence = round(sum(scores) / len(scores), 2) if scores else 0.0
    return OcrCandidate(language, text, confidence)


def _join_ocr_line(parts: tuple[str, ...]) -> str:
    line = " ".join(parts)
    return re.sub(r"(?<=[\u4e00-\u9fff])\s+|\s+(?=[\u4e00-\u9fff])", "", line)


def _confidence_value(value: object) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return -1.0


def _ocr_candidate_rank(candidate: OcrCandidate) -> tuple[float, float, float, int]:
    meaningful = tuple(filter(str.isalpha, candidate.text))
    chinese_count = sum("\u4e00" <= char <= "\u9fff" for char in meaningful)
    latin_count = sum(char.isascii() for char in meaningful)
    total = len(meaningful) or 1
    chinese_ratio = chinese_count / total
    if chinese_count >= 2 and chinese_ratio >= 0.25:
        return 2.0, chinese_ratio, candidate.confidence, len(candidate.text)
    if candidate.language == "eng":
        return 1.0, latin_count / total, candidate.confidence, len(candidate.text)
    return 0.0, chinese_ratio, candidate.confidence, len(candidate.text)


def _configure_tesseract() -> str:
    explicit = os.getenv("TESSERACT_CMD", "").strip()
    command = Path(explicit) if explicit else PROJECT_TESSERACT
    if command.is_file():
        pytesseract.pytesseract.tesseract_cmd = str(command)
        tessdata = command.parent / "tessdata"
        if tessdata.is_dir():
            os.environ["TESSDATA_PREFIX"] = str(tessdata)
        return ""
    discovered = shutil.which("tesseract")
    if discovered:
        pytesseract.pytesseract.tesseract_cmd = discovered
    return ""


def _ocr_warnings(confidence: float | None) -> list[str]:
    if confidence is not None and confidence < LOW_OCR_CONFIDENCE:
        return [f"OCR 平均置信度 {confidence:.1f}%，请逐项核对原图与识别结果"]
    return []


def _rows_to_text(rows: list[dict]) -> str:
    if not rows:
        raise ValueError("文件没有订单明细")
    if len(rows) > MAX_ORDER_ROWS:
        raise ValueError(f"订单文件最多支持 {MAX_ORDER_ROWS} 行商品")
    headers = set(rows[0])
    missing = [header for header in REQUIRED_HEADERS if header not in headers]
    if missing:
        raise ValueError(f"文件缺少列：{'、'.join(missing)}")
    lines = []
    for row in rows:
        product = _cell(row.get("商品"))
        quantity = _number_cell(row.get("数量"))
        unit = _cell(row.get("单位"))
        note = _cell(row.get("备注"))
        if not any((product, quantity, unit, note)):
            continue
        lines.append(f"{product}{quantity}{unit}{note}")
    if not lines:
        raise ValueError("文件没有有效订单明细")
    return "\n".join(lines)


def _cell(value) -> str:
    return "" if value is None else str(value).strip()


def _number_cell(value) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _cell(value)


def _bounded_text(text: str) -> str:
    if len(text) > MAX_NORMALIZED_TEXT_LENGTH:
        raise ValueError(
            f"订单文本不能超过 {MAX_NORMALIZED_TEXT_LENGTH} 个字符"
        )
    return text


def _validate_image_dimensions(width: int, height: int) -> None:
    if (
        width <= 0
        or height <= 0
        or width > MAX_IMAGE_EDGE
        or height > MAX_IMAGE_EDGE
        or width * height > MAX_IMAGE_PIXELS
    ):
        raise ValueError("图片或 PDF 页面像素尺寸过大")
